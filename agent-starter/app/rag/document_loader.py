import logging
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pypdf import PdfReader

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


@dataclass(frozen=True)
class Page:
    """A unit of extracted text. `page` is only set for PDFs (1-based)."""

    text: str
    source: str
    type: str
    page: int | None = None


@dataclass(frozen=True)
class Chunk:
    id: str
    text: str
    source: str
    type: str
    chunk: int
    page: int | None = None

    @property
    def metadata(self) -> dict[str, str | int]:
        meta: dict[str, str | int] = {
            "source": self.source,
            "type": self.type,
            "chunk": self.chunk,
        }
        if self.page is not None:
            meta["page"] = self.page
        return meta


def find_documents(directory: Path) -> list[Path]:
    if not directory.is_dir():
        return []
    return sorted(
        p for p in directory.rglob("*") if p.is_file() and p.suffix.lower() in SUPPORTED_EXTENSIONS
    )


def _load_pdf(path: Path) -> list[Page]:
    reader = PdfReader(str(path))
    return [
        Page(page.extract_text() or "", path.name, "pdf", number)
        for number, page in enumerate(reader.pages, start=1)
    ]


def _load_docx(path: Path) -> list[Page]:
    document = Document(str(path))
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text for cell in row.cells))
    return [Page("\n".join(lines), path.name, "docx")]


def _load_text(path: Path) -> list[Page]:
    text = path.read_text(encoding="utf-8", errors="replace")
    return [Page(text, path.name, path.suffix.lower().lstrip("."))]


_LOADERS = {".pdf": _load_pdf, ".docx": _load_docx, ".txt": _load_text, ".md": _load_text}


def load_document(path: Path) -> list[Page]:
    """Extract text from one file. Documents are only read, never executed."""
    pages = _LOADERS[path.suffix.lower()](path)
    return [page for page in pages if page.text.strip()]


def split_pages(pages: list[Page], chunk_size: int, chunk_overlap: int) -> list[Chunk]:
    splitter = RecursiveCharacterTextSplitter(chunk_size=chunk_size, chunk_overlap=chunk_overlap)
    chunks: list[Chunk] = []
    counters: dict[str, int] = {}
    for page in pages:
        for text in splitter.split_text(page.text):
            index = counters.get(page.source, 0)
            counters[page.source] = index + 1
            chunks.append(
                Chunk(f"{page.source}::{index}", text, page.source, page.type, index, page.page)
            )
    return chunks
